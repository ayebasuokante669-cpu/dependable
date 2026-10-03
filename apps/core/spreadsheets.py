"""The openpyxl boundary: writing an import template, reading a filled one.

Everything in this module knows about spreadsheets and nothing about any
particular entity. It takes a :class:`~apps.core.imports.Sheet` -- the column
contract -- and hands back either ``.xlsx`` bytes or a list of
:class:`~apps.core.imports.SourceRow`: plain strings, one per data row, keyed by
field.

The one rule worth stating: a file that cannot be understood raises
:class:`WorkbookError` with a sentence a school administrator can act on. A
principal who uploads last year's PDF renamed to ``.xlsx`` should be told that,
not shown a stack trace, and a 500 here would be the second time today the
software failed them.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .imports import GUIDANCE_MARKER, Sheet, SourceRow

#: How far down to look for the header row. A school's own export often has a
#: title and a blank line above the real headers; nobody has ten.
HEADER_SEARCH_DEPTH = 12

#: Guards against a sheet whose used range runs to row 1,048,576 because
#: someone once formatted the whole column.
_SCAN_MULTIPLIER = 20


class WorkbookError(Exception):
    """The file could not be read. ``str(exc)`` is shown to the user as-is."""


@dataclass(frozen=True)
class ParsedSheet:
    rows: list[SourceRow]
    sheet_name: str
    #: Headers found in the file that we have no field for -- reported, not
    #: refused, because a school's own sheet carries columns we do not want.
    unknown_columns: list[str]
    #: Optional columns the file did not have at all.
    absent_columns: list[str]


@dataclass(frozen=True)
class Choice:
    """A drop-down down one column, from a short list written into the file."""

    column: str
    options: tuple[str, ...]
    #: Shown by Excel when something else is typed. The field's own label is
    #: used as the title.
    error: str = ""


@dataclass(frozen=True)
class Reference:
    """A list on a sheet of its own, optionally with a drop-down pointing at it.

    The branch's real classes, for the columns that have to name one. Unknown
    class names are the error these imports produce most, and the cheapest place
    to prevent one is before it is typed.
    """

    name: str
    title: str
    values: tuple[str, ...]
    #: The column the drop-down is attached to. Empty means the sheet is a
    #: reference list only -- which is what a column holding *several* class
    #: names needs, because no single-select drop-down can express that.
    column: str = ""
    error: str = ""
    error_title: str = ""
    width: int = 52
    #: A hard stop rather than a warning. Off by default: a school pasting a
    #: column of 300 names should not have Excel refuse every one of them.
    hard: bool = False


# ---------------------------------------------------------------------------
# Writing the template
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="0F2547")   # brand-900
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_GUIDE_FILL = PatternFill("solid", fgColor="EEF1F6")    # ink-100
_GUIDE_FONT = Font(italic=True, color="5A6B85", size=10)
_REQUIRED_FONT = Font(bold=True, color="FFFFFF", size=11)
_THIN = Side(style="thin", color="DFE4ED")


def build_template(
    sheet: Sheet,
    *,
    choices: tuple[Choice, ...] = (),
    references: tuple[Reference, ...] = (),
) -> bytes:
    """The blank import sheet, as ``.xlsx`` bytes.

    Two rows before the school types anything: the headers, and a guidance row
    showing the shape of every field. The guidance row is written so that
    leaving it in place is harmless -- the reader recognises and skips it --
    because half the people who download this will not delete it.
    """
    workbook = Workbook()
    page = workbook.active
    page.title = sheet.name

    for index, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(index)

        header = page.cell(row=1, column=index)
        header.value = f"{column.label} *" if column.required else column.label
        header.fill = _HEADER_FILL
        header.font = _REQUIRED_FONT if column.required else _HEADER_FONT
        header.alignment = Alignment(vertical="center", wrap_text=True)
        header.border = Border(bottom=_THIN)

        guide = page.cell(row=2, column=index)
        guide.value = sheet.guidance_for(column)
        guide.fill = _GUIDE_FILL
        guide.font = _GUIDE_FONT
        guide.alignment = Alignment(vertical="center", wrap_text=True)

        page.column_dimensions[letter].width = column.width

    page.row_dimensions[1].height = 30
    page.row_dimensions[2].height = 30
    # Headers stay put while the school scrolls to row 300.
    page.freeze_panes = "A2"

    for choice in choices:
        _add_choice_list(sheet, page, choice)
    for reference in references:
        _add_reference(sheet, workbook, page, reference)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _column_index(sheet: Sheet, key: str) -> int:
    return next(i for i, c in enumerate(sheet.columns, start=1) if c.key == key)


def _data_range(sheet: Sheet, key: str) -> str:
    letter = get_column_letter(_column_index(sheet, key))
    return f"{letter}3:{letter}{sheet.max_rows + 2}"


def _add_choice_list(sheet: Sheet, page, choice: Choice) -> None:
    """A drop-down down the whole of one column, from row 3 onward."""
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(choice.options) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    validation.error = choice.error or (
        f"Choose one of: {', '.join(choice.options)}."
    )
    validation.errorTitle = sheet.label_for(choice.column)
    page.add_data_validation(validation)
    validation.add(_data_range(sheet, choice.column))


def _add_reference(sheet: Sheet, workbook, page, reference: Reference) -> None:
    """Write a list on its own sheet, and point a drop-down at it if asked."""
    target = workbook.create_sheet(reference.name)
    title = target.cell(row=1, column=1)
    title.value = reference.title
    title.font = Font(bold=True, color="0F2547")
    target.column_dimensions["A"].width = reference.width
    for offset, value in enumerate(reference.values, start=2):
        target.cell(row=offset, column=1).value = value
    target.sheet_state = "visible"

    if not reference.column or not reference.values:
        return

    last = len(reference.values) + 1
    validation = DataValidation(
        type="list",
        formula1=f"='{reference.name}'!$A$2:$A${last}",
        allow_blank=True,
        showDropDown=False,
    )
    validation.error = reference.error
    validation.errorTitle = reference.error_title or sheet.label_for(reference.column)
    validation.errorStyle = "stop" if reference.hard else "warning"
    page.add_data_validation(validation)
    validation.add(_data_range(sheet, reference.column))


# ---------------------------------------------------------------------------
# Reading a filled one
# ---------------------------------------------------------------------------


def read_rows(uploaded, sheet: Sheet) -> ParsedSheet:
    """Parse an uploaded ``.xlsx`` into rows, or raise :class:`WorkbookError`."""
    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises half a dozen different types
        raise WorkbookError(
            "That file could not be opened as an Excel workbook. Save it as "
            ".xlsx (Excel Workbook) and upload it again — .xls, .csv and PDF "
            "files cannot be read here."
        ) from exc

    try:
        page = _first_sheet_with_content(workbook, sheet)
        header_row, positions, unknown = _find_headers(page, sheet)
        rows = _read_data_rows(page, sheet, header_row, positions)
        sheet_name = page.title
    finally:
        workbook.close()

    return ParsedSheet(
        rows=rows,
        sheet_name=sheet_name,
        unknown_columns=unknown,
        absent_columns=[
            c.label for c in sheet.columns
            if c.key not in positions and not c.required
        ],
    )


def _first_sheet_with_content(workbook, sheet: Sheet):
    """The sheet to read: the one named by the contract, else the first with rows.

    Our own templates carry reference sheets -- the branch's class list -- so
    "the first sheet" is not always right once a school has reordered the tabs.
    The contract's own name wins when the file still has it.
    """
    named = workbook.sheetnames
    if sheet.name in named:
        return workbook[sheet.name]
    for page in workbook.worksheets:
        # ``max_row`` is None when the file does not declare its dimensions,
        # which means "unknown", not "empty" -- read it and find out.
        if page.max_row is None or page.max_row >= 1:
            return page
    raise WorkbookError(
        "That workbook is empty — there is no sheet in it with anything on it."
    )


def _find_headers(page, sheet: Sheet) -> tuple[int, dict[str, int], list[str]]:
    """Locate the header row and map each recognised header to its column.

    Matching is by name, so a school that reordered the columns, renamed
    "Surname" to "Last Name", or kept a title line above the headers still
    imports. Position is never assumed.
    """
    best_row = 0
    best: dict[str, int] = {}
    unknown: list[str] = []

    for row_index, row in enumerate(
        page.iter_rows(min_row=1, max_row=HEADER_SEARCH_DEPTH, values_only=True),
        start=1,
    ):
        found: dict[str, int] = {}
        seen_unknown: list[str] = []
        for column_index, value in enumerate(row, start=1):
            key = sheet.column_for_header(value)
            if key is not None and key not in found:
                found[key] = column_index
            elif key is None and value not in (None, ""):
                seen_unknown.append(str(value).strip())
        if len(found) > len(best):
            best_row, best, unknown = row_index, found, seen_unknown

    if len(best) < sheet.header_match_threshold:
        raise WorkbookError(
            "No column headings were recognised in that file. Download the "
            "template, fill it in, and upload it — the first row has to name "
            "the columns."
        )

    missing = [
        sheet.by_key[key].label for key in sheet.required_keys if key not in best
    ]
    if missing:
        raise WorkbookError(
            "That file is missing column"
            f"{'s' if len(missing) > 1 else ''} the import needs: "
            f"{', '.join(missing)}. Download the template and use its headings."
        )

    return best_row, best, sorted(set(unknown))


def _read_data_rows(page, sheet: Sheet, header_row: int, columns: dict[str, int]):
    """Every row below the headers that has something in it."""
    rows: list[SourceRow] = []
    scanned = 0
    limit = sheet.max_rows * _SCAN_MULTIPLIER

    for offset, raw in enumerate(
        page.iter_rows(min_row=header_row + 1, values_only=True), start=1
    ):
        scanned += 1
        if scanned > limit:
            break

        number = header_row + offset
        values = {
            key: _as_text(raw[index - 1]) if index - 1 < len(raw) else ""
            for key, index in columns.items()
        }
        source = SourceRow(number=number, values=values)
        # Blank rows are skipped wherever they are, not just at the end: a
        # spacer between classes is how people lay a roster out by hand.
        if source.is_blank or _is_guidance_row(source, sheet):
            continue

        rows.append(source)
        if len(rows) > sheet.max_rows:
            raise WorkbookError(
                f"That file has more than {sheet.max_rows:,} rows in it. Split "
                f"it into smaller files and import them one at a time."
            )

    if not rows:
        raise WorkbookError(
            f"The columns were recognised but there are no {sheet.noun_plural} "
            f"below them. Fill in the rows and upload the file again."
        )
    return rows


def _is_guidance_row(source: SourceRow, sheet: Sheet) -> bool:
    """The template's own example row, left in place by a hurried user."""
    return source.get(sheet.guidance_column).upper().startswith(GUIDANCE_MARKER)


def _as_text(value) -> str:
    """One cell as the string validation will see.

    Dates come back from Excel as ``datetime``; they are written out in ISO so
    the parsed sheet is plain JSON and survives the session round trip between
    the report and the user confirming it. Whole floats lose their ``.0``,
    because a column of years typed as numbers should not read as ``2019.0``
    and a fee of 45000 should not read as ``45000.0``.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value).strip()
