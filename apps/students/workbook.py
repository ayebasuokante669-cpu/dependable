"""The openpyxl boundary: writing the import template, reading a filled one.

Everything in this module knows about spreadsheets and nothing about
validation. It hands :mod:`apps.students.importer` a list of
:class:`~apps.students.importer.SourceRow` -- plain strings, one per data row,
keyed by field -- and takes back nothing.

The one rule worth stating: a file that cannot be understood raises
:class:`WorkbookError` with a sentence a school administrator can act on. A
principal who uploads last year's PDF renamed to ``.xlsx`` should be told that,
not shown a stack trace, and a 500 here would be the second time today the
software failed them.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .importer import (
    COLUMNS,
    COLUMNS_BY_KEY,
    GUIDANCE_COLUMN,
    GUIDANCE_MARKER,
    MAX_ROWS,
    REQUIRED_KEYS,
    SourceRow,
    column_for_header,
    guidance_for,
)
from .models import StudentStatus

SHEET_NAME = "Students"
REFERENCE_SHEET_NAME = "Classes"
TEMPLATE_FILENAME = "dependable-student-import.xlsx"

#: How far down to look for the header row. A school's own export often has a
#: title and a blank line above the real headers; nobody has ten.
HEADER_SEARCH_DEPTH = 12

#: Enough recognised headers to be confident we found the header row rather
#: than a row of data that happens to contain a word we know.
HEADER_MATCH_THRESHOLD = 3

#: Guards against a sheet whose used range runs to row 1,048,576 because
#: someone once formatted the whole column.
_SCAN_LIMIT = MAX_ROWS * 20


class WorkbookError(Exception):
    """The file could not be read. ``str(exc)`` is shown to the user as-is."""


@dataclass(frozen=True)
class ParsedSheet:
    rows: list[SourceRow]
    sheet_name: str
    #: Headers found in the file that we have no field for -- reported, not
    #: refused, because a school's own sheet carries columns we do not want.
    unknown_columns: list[str]
    #: Optional columns the file did not have at all.
    absent_columns: list[str]


# ---------------------------------------------------------------------------
# Writing the template
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="0F2547")   # brand-900
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_GUIDE_FILL = PatternFill("solid", fgColor="EEF1F6")    # ink-100
_GUIDE_FONT = Font(italic=True, color="5A6B85", size=10)
_REQUIRED_FONT = Font(bold=True, color="FFFFFF", size=11)
_THIN = Side(style="thin", color="DFE4ED")


def build_template(class_names: list[str] | None = None) -> bytes:
    """The blank import sheet, as ``.xlsx`` bytes.

    Two rows before the school types anything: the headers, and a guidance row
    showing the shape of every field. The guidance row is written so that
    leaving it in place is harmless -- the reader recognises and skips it --
    because half the people who download this will not delete it.

    ``class_names`` turns the Class column into a dropdown of the branch's real
    classes and lists them on a second sheet. Unknown classes are the error this
    import produces most, and the cheapest place to prevent one is before it is
    typed.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    for index, column in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)

        header = sheet.cell(row=1, column=index)
        header.value = f"{column.label} *" if column.required else column.label
        header.fill = _HEADER_FILL
        header.font = _REQUIRED_FONT if column.required else _HEADER_FONT
        header.alignment = Alignment(vertical="center", wrap_text=True)
        header.border = Border(bottom=_THIN)

        guide = sheet.cell(row=2, column=index)
        guide.value = guidance_for(column)
        guide.fill = _GUIDE_FILL
        guide.font = _GUIDE_FONT
        guide.alignment = Alignment(vertical="center", wrap_text=True)

        sheet.column_dimensions[letter].width = column.width

    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 30
    # Headers stay put while the school scrolls to row 300.
    sheet.freeze_panes = "A2"

    _add_choice_list(sheet, "sex", ["Male", "Female"])
    _add_choice_list(sheet, "status", [label for _, label in StudentStatus.choices])

    if class_names:
        _add_class_reference(workbook, sheet, class_names)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _column_index(key: str) -> int:
    return next(i for i, c in enumerate(COLUMNS, start=1) if c.key == key)


def _add_choice_list(sheet, key: str, options: list[str]) -> None:
    """A dropdown down the whole of one column, from row 3 onward."""
    letter = get_column_letter(_column_index(key))
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(options) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    validation.error = f"Choose one of: {', '.join(options)}."
    validation.errorTitle = COLUMNS_BY_KEY[key].label
    sheet.add_data_validation(validation)
    validation.add(f"{letter}3:{letter}{MAX_ROWS + 2}")


def _add_class_reference(workbook, sheet, class_names: list[str]) -> None:
    """List the branch's classes on their own sheet and point a dropdown at it."""
    reference = workbook.create_sheet(REFERENCE_SHEET_NAME)
    title = reference.cell(row=1, column=1)
    title.value = "Classes at this branch — use one of these in the Class column"
    title.font = Font(bold=True, color="0F2547")
    reference.column_dimensions["A"].width = 52
    for offset, name in enumerate(class_names, start=2):
        reference.cell(row=offset, column=1).value = name
    reference.sheet_state = "visible"

    last = len(class_names) + 1
    letter = get_column_letter(_column_index("school_class"))
    validation = DataValidation(
        type="list",
        formula1=f"={REFERENCE_SHEET_NAME}!$A$2:$A${last}",
        allow_blank=True,
        showDropDown=False,
    )
    validation.error = (
        "That class does not exist at this branch. Pick one from the list, or "
        "set the class up in Dependable first."
    )
    validation.errorTitle = "Unknown class"
    # A warning, not a hard stop: a school pasting a column of 300 class names
    # should not have Excel refuse every one of them.
    validation.errorStyle = "warning"
    sheet.add_data_validation(validation)
    validation.add(f"{letter}3:{letter}{MAX_ROWS + 2}")


# ---------------------------------------------------------------------------
# Reading a filled one
# ---------------------------------------------------------------------------


def read_rows(uploaded) -> ParsedSheet:
    """Parse an uploaded ``.xlsx`` into rows, or raise :class:`WorkbookError`."""
    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises half a dozen different types
        raise WorkbookError(
            "That file could not be opened as an Excel workbook. Save it as "
            ".xlsx (Excel Workbook) and upload it again — .xls, .csv and PDF "
            "files cannot be read here."
        ) from exc

    try:
        sheet = _first_sheet_with_content(workbook)
        header_row, positions, unknown = _find_headers(sheet)
        rows = _read_data_rows(sheet, header_row, positions)
        sheet_name = sheet.title
    finally:
        workbook.close()

    return ParsedSheet(
        rows=rows,
        sheet_name=sheet_name,
        unknown_columns=unknown,
        absent_columns=[
            c.label for c in COLUMNS if c.key not in positions and not c.required
        ],
    )


def _first_sheet_with_content(workbook):
    for sheet in workbook.worksheets:
        # ``max_row`` is None when the file does not declare its dimensions,
        # which means "unknown", not "empty" -- read it and find out.
        if sheet.max_row is None or sheet.max_row >= 1:
            return sheet
    raise WorkbookError(
        "That workbook is empty — there is no sheet in it with anything on it."
    )


def _find_headers(sheet) -> tuple[int, dict[str, int], list[str]]:
    """Locate the header row and map each recognised header to its column.

    Matching is by name, so a school that reordered the columns, renamed
    "Surname" to "Last Name", or kept a title line above the headers still
    imports. Position is never assumed.
    """
    best_row = 0
    best: dict[str, int] = {}
    unknown: list[str] = []

    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=HEADER_SEARCH_DEPTH, values_only=True),
        start=1,
    ):
        found: dict[str, int] = {}
        seen_unknown: list[str] = []
        for column_index, value in enumerate(row, start=1):
            key = column_for_header(value)
            if key is not None and key not in found:
                found[key] = column_index
            elif key is None and value not in (None, ""):
                seen_unknown.append(str(value).strip())
        if len(found) > len(best):
            best_row, best, unknown = row_index, found, seen_unknown

    if len(best) < HEADER_MATCH_THRESHOLD:
        raise WorkbookError(
            "No column headings were recognised in that file. Download the "
            "template, fill it in, and upload it — the first row has to name "
            "the columns."
        )

    missing = [
        COLUMNS_BY_KEY[key].label for key in REQUIRED_KEYS if key not in best
    ]
    if missing:
        raise WorkbookError(
            "That file is missing column"
            f"{'s' if len(missing) > 1 else ''} the import needs: "
            f"{', '.join(missing)}. Download the template and use its headings."
        )

    return best_row, best, sorted(set(unknown))


def _read_data_rows(sheet, header_row: int, columns: dict[str, int]):
    """Every row below the headers that has something in it."""
    rows: list[SourceRow] = []
    scanned = 0

    for offset, raw in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True), start=1
    ):
        scanned += 1
        if scanned > _SCAN_LIMIT:
            break

        number = header_row + offset
        values = {
            key: _as_text(raw[index - 1]) if index - 1 < len(raw) else ""
            for key, index in columns.items()
        }
        source = SourceRow(number=number, values=values)
        # Blank rows are skipped wherever they are, not just at the end: a
        # spacer between classes is how people lay a roster out by hand.
        if source.is_blank or _is_guidance_row(source):
            continue

        rows.append(source)
        if len(rows) > MAX_ROWS:
            raise WorkbookError(
                f"That file has more than {MAX_ROWS:,} students in it. Split it "
                f"into smaller files and import them one at a time."
            )

    if not rows:
        raise WorkbookError(
            "The columns were recognised but there are no students below them. "
            "Fill in one row per student and upload the file again."
        )
    return rows


def _is_guidance_row(source: SourceRow) -> bool:
    """The template's own example row, left in place by a hurried user."""
    return source.get(GUIDANCE_COLUMN).upper().startswith(GUIDANCE_MARKER)


def _as_text(value) -> str:
    """One cell as the string validation will see.

    Dates come back from Excel as ``datetime``; they are written out in ISO so
    the parsed sheet is plain JSON and survives the session round trip between
    the report and the user confirming it. Whole floats lose their ``.0``,
    because a column of years typed as numbers should not read as ``2019.0``.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value).strip()
