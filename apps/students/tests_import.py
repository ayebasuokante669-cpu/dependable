"""Tests for the Excel student import.

The rules that must never regress: a bad row is reported rather than silently
dropped, a bad row does not take the good rows down with it, the write is
all-or-nothing once it starts, imported students are stamped to the branch the
importer chose and no other, and none of it is reachable by a bursar.

The fixture is the one in :mod:`apps.students.tests` -- one school, two
branches, three students already enrolled -- so the cross-branch cases have
something real to collide with.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.utils import IntegrityError
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.academics.models import Class, Level

from . import importer, workbook
from .models import Student, StudentStatus
from .tests import StudentTestCase

XLSX_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def build_xlsx(rows, *, headers=None, guidance_row=False, title_row=False) -> bytes:
    """A workbook shaped like something a school would actually upload.

    ``rows`` are dicts keyed by column key. ``headers`` overrides the order and
    the spelling of the header row, which is how the reordered-columns and
    header-alias cases are exercised.
    """
    keys = headers or [column.key for column in importer.COLUMNS]
    labels = [
        importer.COLUMNS_BY_KEY[key].label if key in importer.COLUMNS_BY_KEY else key
        for key in keys
    ]

    book = Workbook()
    sheet = book.active
    if title_row:
        sheet.append(["Fulfilled Academy — student list 2025/2026"])
        sheet.append([])
    sheet.append(labels)
    if guidance_row:
        # The real thing, marker and all -- a file that left the template's
        # example row in place is the case being exercised.
        sheet.append([
            importer.guidance_for(importer.COLUMNS_BY_KEY[key])
            if key in importer.COLUMNS_BY_KEY else ""
            for key in keys
        ])
    for row in rows:
        sheet.append([row.get(key, "") for key in keys])

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(content: bytes, name: str = "students.xlsx") -> SimpleUploadedFile:
    return SimpleUploadedFile(name, content, content_type=XLSX_TYPE)


def import_row(**overrides) -> dict:
    """One valid spreadsheet row."""
    row = {
        "admission_number": "FA/2026/001",
        "first_name": "Ifeoma",
        "last_name": "Nwankwo",
        "other_names": "Chidinma",
        "sex": "Female",
        "date_of_birth": "2018-05-14",
        "date_admitted": "2026-01-12",
        "status": "Active",
        "school_class": "Primary 1",
        "parent_name": "Mrs. Uche Nwankwo",
        "parent_phone": "08031122334",
        "parent_email": "uche.nwankwo@example.com",
        "address": "5 Bode Thomas Street, Surulere, Lagos",
    }
    row.update(overrides)
    return row


class ImportTestCase(StudentTestCase):
    """Signed in as the North principal, whose branch needs no choosing."""

    def setUp(self):
        self.client.force_login(self.north_principal)

    def validate(self, dicts, branch=None):
        """Validate rows directly, without the file or the HTTP round trip.

        Row numbers start at 2, the way they do in a real sheet with headers
        on row 1.
        """
        rows = [
            importer.SourceRow(number=index, values=values)
            for index, values in enumerate(dicts, start=2)
        ]
        with self.as_user(self.north_principal):
            return importer.validate(rows, branch=branch or self.north)

    def post_file(self, content, **extra):
        return self.client.post(
            reverse("students:student_import"),
            {"upload": upload(content), **extra},
        )

    def review(self):
        return self.client.get(reverse("students:student_import_review"))

    def confirm(self):
        return self.client.post(
            reverse("students:student_import_review"), {"action": "import"}
        )

    def imported(self):
        return Student.all_objects.filter(admission_number__startswith="FA/2026")


class ImportTemplateTests(ImportTestCase):
    def test_the_template_downloads_as_a_workbook(self):
        response = self.client.get(reverse("students:student_import_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"PK"))  # a zip, i.e. xlsx

    def test_it_carries_a_header_row_and_a_guidance_row(self):
        response = self.client.get(reverse("students:student_import_template"))
        sheet = load_workbook(BytesIO(response.content))["Students"]
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("Admission Number *", headers)   # required, marked
        self.assertIn("Other Names", headers)          # optional, not marked
        self.assertEqual(len(headers), len(importer.COLUMNS))

        guidance = [cell.value for cell in sheet[2]]
        self.assertIn("YYYY-MM-DD", " ".join(guidance))
        self.assertIn("Male or Female", guidance)
        # The row says so, so that leaving it in place is harmless.
        self.assertTrue(guidance[0].startswith("EXAMPLE"))

    def test_it_lists_the_branchs_own_classes_for_the_class_column(self):
        response = self.client.get(reverse("students:student_import_template"))
        book = load_workbook(BytesIO(response.content))
        self.assertIn("Classes", book.sheetnames)
        listed = {
            row[0] for row in book["Classes"].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(listed, {"Primary 1", "JSS 1"})  # South's are not there

    def test_the_template_reads_back_through_our_own_parser(self):
        """Our file has to survive our reader -- guidance row skipped, not
        mistaken for a student."""
        response = self.client.get(reverse("students:student_import_template"))
        with self.assertRaises(workbook.WorkbookError) as caught:
            workbook.read_rows(BytesIO(response.content))
        self.assertIn("no students", str(caught.exception))

    def test_a_bursar_cannot_download_it(self):
        self.client.force_login(self.north_bursar)
        response = self.client.get(reverse("students:student_import_template"))
        self.assertEqual(response.status_code, 403)


class ImportValidationTests(ImportTestCase):
    """The row-level rules."""

    def test_a_clean_sheet_passes_every_row(self):
        report = self.validate([
            import_row(),
            import_row(admission_number="FA/2026/002", first_name="Tunde",
                       school_class="JSS 1"),
        ])
        self.assertEqual(report.ready_count, 2)
        self.assertEqual(report.failed_count, 0)
        self.assertEqual(report.ready[0].student.branch, self.north)
        self.assertEqual(report.ready[0].student.school_id, self.alpha.pk)

    def test_validation_writes_nothing(self):
        before = Student.all_objects.count()
        self.validate([import_row(), import_row(admission_number="FA/2025/001")])
        self.assertEqual(Student.all_objects.count(), before)

    # -- required fields -----------------------------------------------------

    def test_a_missing_required_field_names_its_column(self):
        report = self.validate([import_row(parent_phone="", last_name="")])
        row = report.failed[0]
        self.assertEqual(
            {error.column for error in row.errors},
            {"Surname", "Parent / Guardian Phone"},
        )
        self.assertIn("is required", row.errors[0].message)

    def test_the_optional_fields_really_are_optional(self):
        report = self.validate([import_row(
            other_names="", sex="", date_of_birth="", date_admitted="",
            status="", parent_email="", address="",
        )])
        self.assertEqual(report.ready_count, 1)

    # -- admission numbers ---------------------------------------------------

    def test_a_number_already_on_the_roster_is_refused(self):
        report = self.validate([import_row(admission_number="FA/2025/001")])
        error = report.failed[0].errors[0]
        self.assertEqual(error.column, "Admission Number")
        self.assertIn("already belongs to", error.message)
        self.assertIn("Chinaza Okonkwo", error.message)

    def test_the_clash_is_case_insensitive_like_the_constraint(self):
        report = self.validate([import_row(admission_number="fa/2025/001")])
        self.assertEqual(report.failed_count, 1)

    def test_the_same_number_twice_in_one_file_blames_the_second_row(self):
        report = self.validate([
            import_row(),
            import_row(first_name="Someone", last_name="Else"),
        ])
        self.assertEqual(report.ready_count, 1)
        self.assertEqual(report.ready[0].number, 2)
        self.assertEqual(report.failed[0].number, 3)
        self.assertIn("row 2 of this file", report.failed[0].errors[0].message)

    def test_a_number_in_use_at_another_branch_only_is_accepted(self):
        """Uniqueness is per branch, and the import must not widen it."""
        Student.all_objects.create(
            school_class=self.south_p1, admission_number="FA/2026/001",
            first_name="Somebody", last_name="Southside", sex="male",
            parent_name="A Parent", parent_phone="08034129876",
        )
        self.assertEqual(self.validate([import_row()]).ready_count, 1)

    # -- classes -------------------------------------------------------------

    def test_an_unknown_class_is_refused_with_the_branch_named(self):
        report = self.validate([import_row(school_class="JS1")])
        error = report.failed[0].errors[0]
        self.assertEqual(error.column, "Class")
        self.assertIn("'JS1' is not a class at North", error.message)

    def test_a_class_at_another_branch_does_not_count(self):
        Class.all_objects.create(
            branch=self.south, name="JSS 4", level=Level.JUNIOR_SECONDARY,
            year_in_level=4,
        )
        report = self.validate([import_row(school_class="JSS 4")])
        self.assertIn("not a class at North", report.failed[0].errors[0].message)

    def test_an_inactive_class_says_so_rather_than_not_found(self):
        Class.all_objects.filter(pk=self.north_jss1.pk).update(is_active=False)
        report = self.validate([import_row(school_class="JSS 1")])
        self.assertIn("no longer active", report.failed[0].errors[0].message)

    def test_an_ambiguous_class_name_asks_for_the_arm(self):
        for stream in ("A", "B"):
            Class.all_objects.create(
                branch=self.north, name="JSS 3", stream=stream,
                level=Level.JUNIOR_SECONDARY, year_in_level=3,
            )
        message = self.validate(
            [import_row(school_class="JSS 3")]
        ).failed[0].errors[0].message
        self.assertIn("matches more than one class", message)
        self.assertIn("JSS 3A", message)
        # Naming the arm resolves it.
        self.assertEqual(
            self.validate([import_row(school_class="JSS 3B")]).ready_count, 1
        )

    def test_class_names_are_matched_however_they_are_spaced_and_cased(self):
        for spelling in ("primary 1", "  Primary   1 ", "PRIMARY 1"):
            report = self.validate([import_row(school_class=spelling)])
            self.assertEqual(report.ready_count, 1, spelling)
            self.assertEqual(report.ready[0].student.school_class, self.north_p1)

    # -- everything else -----------------------------------------------------

    def test_a_bad_sex_is_rejected_and_a_blank_one_is_not(self):
        report = self.validate([
            import_row(sex="yes"),
            import_row(admission_number="FA/2026/002", sex=""),
        ])
        self.assertIn("is not a sex", report.failed[0].errors[0].message)
        self.assertEqual(report.ready_count, 1)

    def test_sex_is_accepted_in_the_shapes_people_type_it(self):
        for text in ("M", "male", "Boy", "FEMALE", "f"):
            self.assertEqual(
                self.validate([import_row(sex=text)]).ready_count, 1, text
            )

    def test_dates_are_read_from_the_formats_a_spreadsheet_writes(self):
        for text in ("2018-05-14", "14/05/2018", "14-05-2018", "14 May 2018"):
            report = self.validate([import_row(date_of_birth=text)])
            self.assertEqual(report.ready_count, 1, text)
            self.assertEqual(
                report.ready[0].student.date_of_birth, date(2018, 5, 14), text
            )

    def test_an_unreadable_date_is_a_row_error_not_a_crash(self):
        report = self.validate([import_row(date_of_birth="last April")])
        error = report.failed[0].errors[0]
        self.assertEqual(error.column, "Date of Birth")
        self.assertIn("could not be read as a date", error.message)

    def test_a_missing_admission_date_becomes_today(self):
        report = self.validate([import_row(date_admitted="")])
        self.assertEqual(report.ready[0].student.date_admitted, timezone.localdate())

    def test_the_models_own_cross_field_rules_still_apply(self):
        """Born after being admitted is Student.clean()'s job, not restated here."""
        report = self.validate([
            import_row(date_of_birth="2026-06-01", date_admitted="2026-01-12")
        ])
        self.assertEqual(report.failed_count, 1)

    def test_a_bad_phone_number_is_reported_against_its_column(self):
        report = self.validate([import_row(parent_phone="12345")])
        self.assertEqual(report.failed[0].errors[0].column, "Parent / Guardian Phone")

    def test_excel_eating_the_leading_zero_is_repaired_not_reported(self):
        """A phone column typed as numbers comes back as 8031122334."""
        report = self.validate([import_row(parent_phone="8031122334")])
        self.assertEqual(report.ready_count, 1)
        self.assertEqual(report.ready[0].student.parent_phone, "08031122334")

    def test_a_bad_email_is_reported_and_a_blank_one_is_not(self):
        report = self.validate([
            import_row(parent_email="not an address"),
            import_row(admission_number="FA/2026/002", parent_email=""),
        ])
        self.assertEqual(report.failed[0].errors[0].column, "Parent / Guardian Email")
        self.assertEqual(report.ready_count, 1)

    def test_an_unknown_status_is_refused_and_a_blank_one_means_active(self):
        self.assertIn(
            "is not a status",
            self.validate([import_row(status="paid up")]).failed[0].errors[0].message,
        )
        report = self.validate([import_row(status="")])
        self.assertEqual(report.ready[0].student.status, StudentStatus.ACTIVE)

    def test_a_row_reports_every_one_of_its_problems_at_once(self):
        """The point of the report: one pass through the spreadsheet, not four."""
        report = self.validate([
            import_row(admission_number="FA/2025/001", school_class="JS1",
                       parent_phone="12345", sex="?")
        ])
        row = report.failed[0]
        self.assertEqual(
            {error.column for error in row.errors},
            {"Admission Number", "Class", "Sex", "Parent / Guardian Phone"},
        )
        self.assertTrue(row.summary.startswith("Row 2:"))


class ImportUploadTests(ImportTestCase):
    """The three screens, end to end."""

    def test_the_upload_screen_explains_the_columns(self):
        response = self.client.get(reverse("students:student_import"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Admission Number")
        self.assertContains(response, "Download template")

    def test_a_clean_file_imports_every_row(self):
        content = build_xlsx([
            import_row(),
            import_row(admission_number="FA/2026/002", first_name="Tunde",
                       last_name="Bakare", school_class="JSS 1"),
            import_row(admission_number="FA/2026/003", first_name="Halima",
                       last_name="Sule"),
        ], guidance_row=True)

        self.assertRedirects(
            self.post_file(content), reverse("students:student_import_review")
        )
        report = self.review().context["report"]
        self.assertEqual((report.total, report.ready_count, report.failed_count),
                         (3, 3, 0))
        self.assertEqual(self.imported().count(), 0)  # still nothing written

        self.assertRedirects(self.confirm(), reverse("students:student_list"))
        imported = list(self.imported())
        self.assertEqual(len(imported), 3)
        self.assertEqual({s.branch_id for s in imported}, {self.north.pk})
        self.assertEqual({s.school_id for s in imported}, {self.alpha.pk})
        # The pending file is gone, so a refresh cannot import it twice.
        self.assertNotIn("students.import", self.client.session)

    def test_a_file_with_bad_rows_imports_the_good_ones(self):
        """The whole point: partial, with the failures listed for correction."""
        content = build_xlsx([
            import_row(),                                           # fine
            import_row(admission_number="FA/2025/001"),             # already taken
            import_row(admission_number="FA/2026/003", school_class="JS1"),
            import_row(admission_number="FA/2026/004", first_name="Ada",
                       last_name="Eze", school_class="JSS 1"),      # fine
            import_row(admission_number="FA/2026/005", parent_phone=""),
        ])
        self.post_file(content)

        report = self.review().context["report"]
        self.assertEqual((report.total, report.ready_count, report.failed_count),
                         (5, 2, 3))
        self.assertEqual([row.number for row in report.failed], [3, 4, 6])

        self.confirm()
        self.assertEqual(
            set(self.imported().values_list("admission_number", flat=True)),
            {"FA/2026/001", "FA/2026/004"},
        )

    def test_the_report_names_the_field_and_the_reason(self):
        self.post_file(build_xlsx([
            import_row(admission_number="FA/2025/001", school_class="JS1")
        ]))
        response = self.review()
        self.assertContains(response, "Row 2")
        self.assertContains(response, "already belongs to")
        self.assertContains(response, "is not a class at North")

    def test_a_file_of_only_bad_rows_imports_nothing_and_stays_put(self):
        self.post_file(build_xlsx([import_row(school_class="Nowhere")]))
        response = self.confirm()
        self.assertEqual(response.status_code, 200)  # back on the report
        self.assertEqual(self.imported().count(), 0)

    def test_a_number_duplicated_inside_the_file_imports_once(self):
        self.post_file(build_xlsx([
            import_row(),
            import_row(first_name="Different", last_name="Child"),
        ]))
        self.confirm()
        self.assertEqual(self.imported().count(), 1)

    def test_a_reordered_file_with_the_schools_own_headings_still_imports(self):
        """The realistic upload: their spreadsheet, not ours."""
        content = build_xlsx(
            [
                {
                    "adm no": " FA/2026/010 ", "given name": "Nkem",
                    "surname": "  Okoro  ", "gender": "male", "dob": "2018-05-14",
                    "class": "primary 1", "guardian": "Mr. Paul Okoro",
                    "phone": "0803 112 2334",
                    "house": "Blue",  # a column of theirs we have no field for
                },
                {},  # a spacer row between classes
                {
                    "adm no": "FA/2026/011", "given name": "Amina",
                    "surname": "Bello", "gender": "female", "dob": "",
                    "class": "JSS 1", "guardian": "Mrs. Sadiq Bello",
                    "phone": "08031122335", "house": "Red",
                },
                {},  # trailing blanks
                {},
            ],
            headers=["adm no", "given name", "surname", "gender", "dob", "class",
                     "guardian", "phone", "house"],
            title_row=True,
        )
        self.post_file(content)

        report = self.review().context["report"]
        self.assertEqual((report.total, report.ready_count), (2, 2))
        self.assertEqual(report.unknown_columns, ["house"])
        self.assertIn("Address", report.absent_columns)

        self.confirm()
        student = Student.all_objects.get(admission_number="FA/2026/010")
        self.assertEqual(student.last_name, "Okoro")           # whitespace gone
        self.assertEqual(student.parent_phone, "08031122334")  # spacing gone
        self.assertEqual(student.school_class, self.north_p1)
        self.assertEqual(student.status, StudentStatus.ACTIVE)  # column absent

    # -- files that are not what they should be ------------------------------

    def test_a_file_that_is_not_a_workbook_is_a_form_error(self):
        response = self.post_file(b"this is not a spreadsheet at all")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "could not be opened as an Excel workbook")
        self.assertNotIn("students.import", self.client.session)

    def test_a_workbook_missing_a_required_column_says_which(self):
        content = build_xlsx(
            [{"first_name": "Ada", "last_name": "Eze", "sex": "Female"}],
            headers=["first_name", "last_name", "sex"],
        )
        response = self.post_file(content)
        self.assertContains(response, "missing column")
        self.assertContains(response, "Admission Number")

    def test_a_workbook_with_nothing_recognisable_is_refused_gracefully(self):
        content = build_xlsx(
            [{}], headers=["invoice ref", "amount", "paid on"]
        )
        response = self.post_file(content)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "No column headings were recognised")

    def test_an_untouched_template_is_refused_gracefully(self):
        response = self.post_file(build_xlsx([], guidance_row=True))
        self.assertContains(response, "no students below them")

    def test_a_non_xlsx_extension_is_refused_before_it_is_read(self):
        response = self.client.post(
            reverse("students:student_import"),
            {"upload": upload(b"name,class\nAda,JSS 1\n", name="roster.csv")},
        )
        self.assertContains(response, "not an .xlsx file")

    # -- the pending file ----------------------------------------------------

    def test_arriving_at_the_report_with_nothing_pending_goes_back(self):
        self.assertRedirects(self.review(), reverse("students:student_import"))

    def test_discarding_leaves_nothing_behind(self):
        self.post_file(build_xlsx([import_row()]))
        response = self.client.post(
            reverse("students:student_import_review"), {"action": "discard"}
        )
        self.assertRedirects(response, reverse("students:student_import"))
        self.assertNotIn("students.import", self.client.session)
        self.assertEqual(self.imported().count(), 0)

    def test_the_file_is_rechecked_against_the_roster_before_it_is_written(self):
        """The number was free when the file was uploaded and taken by the time
        the principal pressed the button."""
        self.post_file(build_xlsx([import_row()]))
        self.assertEqual(self.review().context["report"].ready_count, 1)

        Student.all_objects.create(
            school_class=self.north_p1, admission_number="FA/2026/001",
            first_name="Got", last_name="There First", sex="male",
            parent_name="A Parent", parent_phone="08034129876",
        )
        response = self.confirm()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.imported().count(), 1)  # only the one already there

    def test_a_failure_partway_through_writes_nothing(self):
        """One transaction: a batch that dies on the third student must not
        leave the first two behind."""
        self.post_file(build_xlsx([
            import_row(),
            import_row(admission_number="FA/2026/002"),
            import_row(admission_number="FA/2026/003"),
        ]))

        real_save, calls = Student.save, {"n": 0}

        def explode(self, *args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 3:
                raise IntegrityError("a constraint fired on the last one")
            return real_save(self, *args, **kwargs)

        with patch.object(Student, "save", explode):
            response = self.confirm()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.imported().count(), 0)


class ImportPermissionTests(ImportTestCase):
    IMPORT_URLS = (
        "students:student_import",
        "students:student_import_template",
        "students:student_import_review",
    )

    def test_a_bursar_is_refused_every_import_screen(self):
        self.client.force_login(self.north_bursar)
        for name in self.IMPORT_URLS:
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)

    def test_a_bursar_cannot_post_a_file_either(self):
        self.client.force_login(self.north_bursar)
        response = self.post_file(build_xlsx([import_row()]))
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.imported().count(), 0)

    def test_the_roster_offers_the_import_only_to_those_who_may_use_it(self):
        response = self.client.get(reverse("students:student_list"))
        self.assertContains(response, reverse("students:student_import"))

        self.client.force_login(self.north_bursar)
        response = self.client.get(reverse("students:student_list"))
        self.assertNotContains(response, reverse("students:student_import"))

    def test_signed_out_visitors_are_sent_to_the_login_page(self):
        self.client.logout()
        response = self.client.get(reverse("students:student_import"))
        self.assertEqual(response.status_code, 302)
        self.assertIn("login", response["Location"])

    def test_a_principal_is_not_asked_which_campus(self):
        response = self.client.get(reverse("students:student_import"))
        self.assertNotIn("branch", response.context["form"].fields)

    def test_a_principal_cannot_import_into_another_branch(self):
        """The field does not exist for them, and a forged one is ignored."""
        self.post_file(build_xlsx([import_row()]), branch=self.south.pk)
        self.confirm()
        student = Student.all_objects.get(admission_number="FA/2026/001")
        self.assertEqual(student.branch, self.north)
        self.assertEqual(student.school_class, self.north_p1)

    def test_a_school_owner_picks_the_campus_and_it_is_honoured(self):
        self.client.force_login(self.alpha_owner)
        response = self.client.get(reverse("students:student_import"))
        self.assertIn("branch", response.context["form"].fields)

        # "Primary 1" runs at both campuses; the chosen branch decides which.
        self.post_file(build_xlsx([import_row()]), branch=self.south.pk)
        self.confirm()
        student = Student.all_objects.get(admission_number="FA/2026/001")
        self.assertEqual(student.branch, self.south)
        self.assertEqual(student.school_class, self.south_p1)

    def test_a_school_owner_is_offered_a_template_per_campus(self):
        """One class list per branch, so the drop-down is never the wrong one."""
        self.client.force_login(self.alpha_owner)
        links = self.client.get(reverse("students:student_import")).context[
            "template_links"
        ]
        self.assertEqual(
            [link["label"] for link in links],
            ["Download template — North", "Download template — South"],
        )

    def test_a_school_owners_template_follows_the_branch_they_asked_for(self):
        self.client.force_login(self.alpha_owner)
        response = self.client.get(
            reverse("students:student_import_template"), {"branch": self.south.pk}
        )
        book = load_workbook(BytesIO(response.content))
        listed = {
            row[0] for row in book["Classes"].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(listed, {"Primary 1"})  # South runs one class
